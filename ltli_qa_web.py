import io
from pathlib import Path
import re
import time
from collections import Counter

import pandas as pd
import streamlit as st
from docx import Document

from telemetry import (
    APP_VERSION, elapsed_milliseconds, file_size_bucket,
    telemetry_enabled, track_event, track_once,
)


st.set_page_config(
    page_title="LTLI QA",
    page_icon="🔎",
    layout="wide",
    initial_sidebar_state="expanded",
)

LANGUAGE = {
    "ko": {
        "hero_title": "반복적인 번역 검사를<br>더 간편하게.", "hero_text": "숫자, 태그, 자리표시자, 용어 등 번역문에서 발생하기 쉬운 문제를 자동으로 점검합니다. 검사 결과를 확인하고 Excel 보고서로 다운로드하세요.",
        "brand_name": "LTLI · 언어와 기술의 연결", "public_beta": "LTLI QA", "free": "무료 이용", "formats": "XLSX · DOCX · TXT", "no_account": "회원가입 불필요",
        "step1": "파일 추가", "step1_text": "이중언어 XLSX 또는 정렬된 파일 쌍을 업로드합니다.", "step2": "15개 검사 실행", "step2_text": "숫자, 태그, 용어, 공백 등을 확인합니다.", "step3": "결과 내보내기", "step3_text": "공유 가능한 Excel QA 보고서를 다운로드합니다.",
        "policy": "🔒 업로드 전 확인 — 파일 업로드 및 데이터 보호 안내", "policy_text": "공개 가능한 자료만 업로드하시기 바랍니다. 개인정보, 고객 기밀, 외교문서 또는 미공개 상업 자료를 업로드하지 마십시오. 파일당 최대 크기는 10MB입니다.",
        "feedback": "베타 피드백", "feedback_intro": "여러분의 의견은 LTLI QA 개선에 활용됩니다. 언제든 제출할 수 있습니다.", "feedback_callout": "💬 QA를 사용해 보셨나요? 2분 피드백을 남겨 주세요.", "feedback_after_qa": "💬 결과는 어떠셨나요? 베타 피드백 남기기", "feedback_after_qa_text": "여러분의 실제 사용 경험은 검사 정확도와 다음 기능의 우선순위를 정하는 데 큰 도움이 됩니다.", "satisfaction": "전체 만족도", "helpful": "결과가 실제 검수에 도움이 되었나요?", "false_positive": "오탐은 어느 정도였나요?", "useful": "가장 유용한 검사", "features": "추가로 필요한 기능", "reuse": "다시 사용할 의향이 있나요?", "role": "직업군(선택)", "comments": "자유 의견(선택)", "interview": "후속 인터뷰 이메일(선택)", "consent": "베타 인터뷰 연락 목적으로 이 이메일을 저장하는 데 동의합니다.", "privacy_caption": "익명 피드백에는 무작위 세션 ID만 사용합니다. 자유 의견에 개인정보·고객 기밀·외교문서·미공개 정보를 입력하지 마십시오. 선택 이메일은 사용 이벤트와 분리해 저장합니다.", "send": "피드백 보내기", "thanks": "감사합니다. 피드백이 제출되었습니다.",
        "add_files": "1. 파일 추가", "choose_format": "번역 작업에 맞는 입력 형식을 선택하세요.", "xlsx_tab": "📊 이중언어 XLSX", "pair_tab": "📄 DOCX / TXT 파일 쌍", "begin": "위 탭에서 형식을 선택하고 파일을 업로드하면 시작됩니다.", "drop_xlsx": "이중언어 스프레드시트를 업로드하세요", "required_columns": "필수 열: Source 및 Target · 최대 10MB", "source_file": "원문 파일", "target_file": "번역문 파일", "pair_help": "같은 형식의 파일 두 개를 사용하세요. DOCX는 비어 있지 않은 문단, TXT는 비어 있지 않은 줄을 기준으로 정렬합니다.", "glossary": "＋ 용어집 추가(선택)", "glossary_file": "용어집 파일",
        "review": "2. 미리보기 및 QA 실행", "ready": "개의 정렬된 세그먼트가 준비되었습니다. 미리보기를 확인한 후 검사를 시작하세요.", "run": "QA 검사 실행 →", "results": "3. QA 결과", "overview": "#### 오류 요약", "all_segments": "#### 전체 세그먼트", "review_issues": "#### 검토가 필요한 오류", "no_issues": "QA 오류가 발견되지 않았습니다.", "download": "Excel 보고서 다운로드 ↓",
        'contact': '문의 및 제안: contact@ltli.kr · 홈페이지: https://ltli.ai',
        'xlsx_help': '스프레드시트에는 Source 열과 Target 열이 각각 하나씩 있어야 합니다.',
        'glossary_help': '필수 열: Source Term 및 Target Term (공백·밑줄·대소문자 차이 허용).',
        'upload_xlsx': '이중언어 XLSX 파일을 업로드하세요.',
        'upload_pair': '원문 파일과 번역문 파일을 모두 업로드하세요.',
        'current_columns': '현재 열:',
        'qa_failed': 'QA를 완료하지 못했습니다. 입력 파일을 확인한 후 다시 시도하세요.',
        'total': '전체 세그먼트',
        'passed': '통과',
        'errors': '검토 필요',
        'pass_rate': '통과율',
        'format': '파일 형식',
        'source_units': '원문 단위 수',
        'target_units': '번역문 단위 수',
        'glossary_loaded': '용어집 항목 {count}개를 불러왔습니다.',
        'uploaded': '업로드 완료: {name}',
        'file_read_failed': '파일을 읽지 못했습니다. 파일 형식과 크기(최대 10MB)를 확인하세요.',
        'issue_count': '{count}개 세그먼트에서 검토가 필요한 문제가 발견되었습니다.',
        'aligned': '{unit} {count}개를 정렬했습니다.',
        'paragraphs': '문단',
        'lines': '줄',
        'checks_description': '번역 누락, 원문과 동일한 번역, 숫자, 구두점, 태그, 자리표시자, 괄호, URL, 이메일, 공백, 파일명, 단축키, 단위 및 용어를 검사합니다. 자동 QA는 최종 언어 검수를 대체하지 않습니다.',
        'downloads_title': '📥 샘플 파일 및 빈 템플릿 다운로드',
        'downloads_help': 'Excel 샘플은 원문과 번역문이 포함된 하나의 파일입니다. Word 샘플은 원문과 번역문을 각각 다운로드하여 함께 업로드하세요. 샘플에는 가상 데이터와 의도적인 QA 오류가 포함되어 있습니다.',
        'sample_label': '테스트용 샘플 파일',
        'template_label': '작성용 빈 템플릿',
        'download_xlsx': '⬇ Excel (.xlsx)',
        'download_source': '⬇ Word 원문 파일 (.docx)',
        'download_target': '⬇ Word 번역문 파일 (.docx)',
        'downloads_unavailable': '파일이 아직 제공되지 않습니다',
        'template_help': 'Excel 템플릿에는 Source 및 Target 열이 있습니다. Word 템플릿은 별도의 빈 문서입니다. 원문과 번역문의 문단 수와 순서를 맞춰 작성하세요.',
        'validation_0': '파일에는 Source 열과 Target 열이 각각 하나씩 있어야 합니다. 열 이름은 대소문자를 구분하지 않습니다.',
        'validation_1': '{name} 파일이 비어 있습니다.',
        'validation_2': '{name} 파일이 10MB 크기 제한을 초과합니다.',
        'validation_3': 'DOCX 파일을 읽지 못했습니다: {name}. 파일이 손상되지 않았는지 확인하세요.',
        'validation_4': '{name} 파일을 해독하지 못했습니다. UTF-8, UTF-16 또는 CP949로 저장하세요.',
        'validation_5': '원문과 번역문 파일의 형식이 같아야 합니다(DOCX + DOCX 또는 TXT + TXT).',
        'validation_6': 'DOCX 및 TXT 파일 쌍만 지원합니다.',
        'validation_7': '원문에서 비어 있지 않은 세그먼트를 찾지 못했습니다.',
        'validation_8': '번역문에서 비어 있지 않은 세그먼트를 찾지 못했습니다.',
        'validation_9': '문서 구조가 다릅니다. 원문은 {source_count}개, 번역문은 {target_count}개입니다. 비어 있지 않은 문단 또는 줄의 수와 순서를 맞춘 후 다시 실행하세요.',
        'validation_10': '용어집은 XLSX 또는 CSV 파일이어야 합니다.',
        'validation_11': '용어집 파일을 읽지 못했습니다. 파일 형식과 내용을 확인하세요.',
        'validation_12': '용어집에는 Source Term 및 Target Term 열이 필요합니다. source_term 및 target_term 표기도 지원합니다.',
        'footer': '언어기술현지화연구소 · LTLI QA · 번역가, 검수자 및 현지화 팀을 위한 도구',
    },
    "en": {}
}

query_language = str(st.query_params.get("lang", "")).lower()
try:
    browser_language = str(getattr(st.context, "locale", "en")).lower()
except (AttributeError, RuntimeError):
    browser_language = "en"
detected_language = query_language if query_language in {"ko", "en"} else ("ko" if browser_language.startswith("ko") else "en")
if "ui_language" not in st.session_state:
    st.session_state["ui_language"] = detected_language

language_labels = {"한국어": "ko", "English": "en"}
selected_label = st.selectbox("Language / 언어", list(language_labels), index=0 if st.session_state["ui_language"] == "ko" else 1)
selected_language = language_labels[selected_label]
if selected_language != st.session_state["ui_language"]:
    st.session_state["ui_language"] = selected_language
    st.query_params["lang"] = selected_language
    st.rerun()

UI_LANGUAGE = st.session_state["ui_language"]

def L(key, english):
    return LANGUAGE.get(UI_LANGUAGE, {}).get(key, english)

ISSUE_LABELS_KO = {'Missing translation': '번역 누락', 'Source and target are identical': '원문과 번역문 동일', 'Number mismatch': '숫자 불일치', 'Punctuation mismatch': '구두점 불일치', 'Tag mismatch': '태그 불일치', 'Placeholder mismatch': '자리표시자 불일치', 'Bracket mismatch': '괄호 불일치', 'Unbalanced target brackets': '번역문 괄호 짝 불일치', 'URL mismatch': 'URL 불일치', 'Email mismatch': '이메일 불일치', 'Whitespace warning': '공백 주의', 'File name mismatch': '파일명 불일치', 'Shortcut mismatch': '단축키 불일치', 'Unit mismatch': '단위 불일치', 'Terminology mismatch': '용어 불일치', 'Leading space': '앞쪽 공백', 'Trailing space': '뒤쪽 공백', 'Multiple spaces': '연속 공백', 'Tab character': '탭 문자'}

def localized_results(frame):
    """Translate presentation copies; preserve reports, rule IDs, and analytics."""
    if UI_LANGUAGE != "ko":
        return frame
    display = frame.copy()
    def translate(value):
        if not isinstance(value, str):
            return value
        for english, korean in ISSUE_LABELS_KO.items():
            value = value.replace(english, korean)
        return value
    for column in ("QA Issues", "Issue Type"):
        if column in display:
            display[column] = display[column].map(translate)
    if "QA Result" in display:
        display["QA Result"] = display["QA Result"].replace({"Pass": "통과", "Error": "검토 필요"})
    return display.rename(columns={"Segment": "세그먼트", "Source": "원문", "Target": "번역문", "QA Issues": "검출 사항", "QA Result": "검사 결과", "Issue Type": "문제 유형", "Count": "건수"})


hero_html = (
    """
    <style>
    :root {
        --ltli-navy: #10223d;
        --ltli-cobalt: #3157e2;
        --ltli-violet: #7c4dff;
        --ltli-cyan: #25b6d2;
        --ltli-surface: #f6f8ff;
        --ltli-line: #e3e8f5;
    }
    .stApp {
        background:
            radial-gradient(circle at 92% 5%, rgba(124,77,255,.10), transparent 25rem),
            radial-gradient(circle at 3% 25%, rgba(37,182,210,.08), transparent 24rem),
            #fbfcff;
        color: var(--ltli-navy);
    }
    .block-container {max-width: 1180px; padding-top: 2rem; padding-bottom: 4rem;}
    .ltli-hero {
        position: relative; overflow: hidden; padding: 2.25rem 2.4rem;
        border: 1px solid rgba(255,255,255,.65); border-radius: 28px;
        background: linear-gradient(125deg, #142d62 0%, #365ce7 55%, #7c4dff 100%);
        color: white; box-shadow: 0 18px 50px rgba(49,87,226,.20);
        margin-bottom: 1.25rem;
    }
    .ltli-hero:after {
        content: ""; position: absolute; width: 270px; height: 270px;
        border-radius: 50%; right: -70px; top: -120px;
        background: rgba(255,255,255,.12);
    }
    .ltli-brand {font-size: .82rem; font-weight: 750; letter-spacing: .16em; opacity: .78;}
    .ltli-hero h1 {font-size: clamp(2rem, 4vw, 3.35rem); line-height: 1.05; margin: .65rem 0 .75rem; color: white;}
    .ltli-hero p {font-size: 1.05rem; max-width: 690px; opacity: .92; margin: 0;}
    .ltli-badges {display:flex; flex-wrap:wrap; gap:.55rem; margin-top:1.35rem;}
    .ltli-badge {padding:.42rem .72rem; border-radius:999px; background:rgba(255,255,255,.14); border:1px solid rgba(255,255,255,.18); font-size:.78rem;}
    .ltli-steps {display:grid; grid-template-columns:repeat(3,1fr); gap:.75rem; margin:1rem 0 1.7rem;}
    .ltli-step {background:white; border:1px solid var(--ltli-line); border-radius:16px; padding:.85rem 1rem; box-shadow:0 6px 18px rgba(16,34,61,.045);}
    .ltli-step strong {display:block; font-size:.86rem; color:var(--ltli-cobalt); margin-bottom:.15rem;}
    .ltli-step span {font-size:.82rem; color:#63708a;}
    .ltli-section {font-size:1.35rem; font-weight:760; color:var(--ltli-navy); margin:1.6rem 0 .25rem;}
    div[data-testid="stFileUploader"] {background:white; border:1px solid var(--ltli-line); border-radius:18px; padding:1rem 1rem .45rem; box-shadow:0 8px 24px rgba(16,34,61,.05);}
    div[data-testid="stFileUploaderDropzone"] {background:var(--ltli-surface); border:1.5px dashed #b9c5e9; border-radius:14px;}
    div[data-testid="stMetric"] {background:white; border:1px solid var(--ltli-line); border-radius:16px; padding:1rem; box-shadow:0 6px 20px rgba(16,34,61,.05);}
    div[data-testid="stMetricValue"] {color:var(--ltli-cobalt);}
    div[data-baseweb="tab-list"] {gap:.35rem; background:#eef2ff; padding:.3rem; border-radius:14px;}
    button[data-baseweb="tab"] {border-radius:10px;}
    .stButton > button[kind="primary"], .stDownloadButton > button {
        min-height:3rem; border-radius:12px; font-weight:700; border:0;
        box-shadow:0 8px 20px rgba(49,87,226,.18);
    }
    .ltli-footer {text-align:center; color:#748098; font-size:.78rem; margin-top:2.5rem;}
    .ltli-feedback-callout {margin:1rem 0; padding:1rem 1.15rem; border-radius:16px; color:#24376b; font-weight:700; background:linear-gradient(110deg,#eef3ff,#f4edff); border:1px solid #cdd8ff; box-shadow:0 8px 24px rgba(49,87,226,.08);}
    @media (max-width: 700px) {
        .block-container {padding-top:1rem;}
        .ltli-hero {padding:1.6rem 1.3rem; border-radius:22px;}
        .ltli-steps {grid-template-columns:1fr;}
    }
    </style>
    <section class="ltli-hero">
      <div class="ltli-brand">LTLI · Language &amp; Technology</div>
      <h1>Simplify Routine<br>Translation QA.</h1>
      <p>Automatically check translations for issues involving numbers, tags, placeholders, terminology, and more. Review the findings and download an Excel QA report.</p>
      <div class="ltli-badges">
        <span class="ltli-badge">Free to use</span>
        <span class="ltli-badge">XLSX · DOCX · TXT</span><span class="ltli-badge">No account required</span>
      </div>
    </section>
    <div class="ltli-steps">
      <div class="ltli-step"><strong>01 · Add your files</strong><span>Upload bilingual XLSX or an aligned file pair.</span></div>
      <div class="ltli-step"><strong>02 · Run 15 checks</strong><span>Review numbers, tags, terms, spacing, and more.</span></div>
      <div class="ltli-step"><strong>03 · Export results</strong><span>Download a shareable Excel QA report.</span></div>
    </div>
    """,
)
hero_html = hero_html[0]
for source, target in {
    "Simplify Routine<br>Translation QA.": L("hero_title", "Simplify Routine<br>Translation QA."),
    "Automatically check translations for issues involving numbers, tags, placeholders, terminology, and more. Review the findings and download an Excel QA report.": L("hero_text", "Automatically check translations for issues involving numbers, tags, placeholders, terminology, and more. Review the findings and download an Excel QA report."),
    "LTLI · Language &amp; Technology": L("brand_name", "LTLI · Language &amp; Technology"), "Public beta": L("public_beta", "LTLI QA"), "Free to use": L("free", "Free to use"), "No account required": L("no_account", "No account required"),
    "Add your files": L("step1", "Add your files"), "Upload bilingual XLSX or an aligned file pair.": L("step1_text", "Upload bilingual XLSX or an aligned file pair."),
    "Run 15 checks": L("step2", "Run 15 checks"), "Review numbers, tags, terms, spacing, and more.": L("step2_text", "Review numbers, tags, terms, spacing, and more."),
    "Export results": L("step3", "Export results"), "Download a shareable Excel QA report.": L("step3_text", "Download a shareable Excel QA report."),
}.items():
    hero_html = hero_html.replace(source, target)
st.markdown(hero_html, unsafe_allow_html=True)

with st.expander(L("policy", "🔒 Before you upload — file handling policy"), expanded=False):
    st.warning(
        L("policy_text", "Use public test data only. Do not upload personal information, customer-confidential files, diplomatic documents, or unpublished commercial content. Maximum file size: 10 MB.")
    )
    st.caption(L('contact', 'Inquiries & Suggestions: contact@ltli.kr · Website: https://ltli.ai'))

track_once("session_started", "session_started", {"app_version": APP_VERSION})

# --------------------------------------------------
# 공통 함수
# --------------------------------------------------
def normalize_text(text):
    if pd.isna(text):
        return ""

    return str(text).strip()


def normalize_bilingual_columns(dataframe):
    """Return a copy whose source/target columns use canonical names.

    Column matching ignores case and common separators, so both the legacy
    ``Source``/``Target`` format and files using ``source``/``target`` are
    accepted. Other metadata columns are preserved unchanged.
    """
    normalized_dataframe = dataframe.copy()
    normalized_dataframe.columns = [
        str(column).strip().lstrip("\ufeff")
        for column in normalized_dataframe.columns
    ]

    column_lookup = {}
    for column in normalized_dataframe.columns:
        normalized_name = re.sub(r"[\s_-]+", "", column).casefold()
        column_lookup.setdefault(normalized_name, []).append(column)

    source_columns = column_lookup.get("source", [])
    target_columns = column_lookup.get("target", [])

    if len(source_columns) != 1 or len(target_columns) != 1:
        raise ValueError(
            L('validation_0', "The file must contain exactly one Source column and one Target "
            "column. Column names are case-insensitive.")
        )

    return normalized_dataframe.rename(
        columns={
            source_columns[0]: "Source",
            target_columns[0]: "Target",
        }
    )


MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024


def validate_file_size(uploaded_file):
    """Reject empty files and files larger than the MVP upload limit."""
    file_size = getattr(uploaded_file, "size", None)

    if file_size == 0:
        raise ValueError((LANGUAGE["ko"]['validation_1'].format(name=uploaded_file.name) if UI_LANGUAGE == "ko" else f"{uploaded_file.name} is empty."))

    if file_size is not None and file_size > MAX_FILE_SIZE_BYTES:
        raise ValueError(
            (LANGUAGE["ko"]['validation_2'].format(name=uploaded_file.name) if UI_LANGUAGE == "ko" else f"{uploaded_file.name} exceeds the 10 MB file-size limit.")
        )


def get_file_extension(uploaded_file):
    file_name = uploaded_file.name

    if "." not in file_name:
        return ""

    return file_name.rsplit(".", 1)[-1].lower()


def read_docx_segments(uploaded_file):
    """Read non-empty DOCX paragraphs in document order."""
    validate_file_size(uploaded_file)

    try:
        document = Document(io.BytesIO(uploaded_file.getvalue()))
    except Exception as error:
        raise ValueError(
            (LANGUAGE["ko"]['validation_3'].format(name=uploaded_file.name) if UI_LANGUAGE == "ko" else f"Failed to read DOCX file {uploaded_file.name}: {error}")
        ) from error

    return [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]


def decode_text_bytes(raw_bytes, file_name):
    """Decode common Korean and Unicode TXT encodings without guessing silently."""
    if raw_bytes.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16",)
    else:
        encodings = ("utf-8-sig", "utf-8", "cp949")

    for encoding in encodings:
        try:
            return raw_bytes.decode(encoding), encoding
        except (UnicodeDecodeError, UnicodeError):
            continue

    raise ValueError(
        (LANGUAGE["ko"]['validation_4'].format(name=file_name) if UI_LANGUAGE == "ko" else f"Could not decode {file_name}. Save it as UTF-8, UTF-16, or CP949.")
    )


def read_txt_segments(uploaded_file):
    """Read non-empty TXT lines in file order."""
    validate_file_size(uploaded_file)
    decoded_text, encoding = decode_text_bytes(
        uploaded_file.getvalue(),
        uploaded_file.name
    )
    segments = [line.strip() for line in decoded_text.splitlines() if line.strip()]
    return segments, encoding


def build_parallel_dataframe(source_file, target_file):
    """Build Source/Target rows from two structurally matching documents."""
    source_extension = get_file_extension(source_file)
    target_extension = get_file_extension(target_file)

    if source_extension != target_extension:
        raise ValueError(
            L('validation_5', "Source and target files must use the same format (DOCX + DOCX or TXT + TXT).")
        )

    metadata = {"format": source_extension.upper()}

    if source_extension == "docx":
        source_segments = read_docx_segments(source_file)
        target_segments = read_docx_segments(target_file)
        metadata["unit"] = "paragraphs"
    elif source_extension == "txt":
        source_segments, source_encoding = read_txt_segments(source_file)
        target_segments, target_encoding = read_txt_segments(target_file)
        metadata.update(
            {
                "unit": "lines",
                "source_encoding": source_encoding,
                "target_encoding": target_encoding,
            }
        )
    else:
        raise ValueError(L('validation_6', "Only DOCX and TXT file pairs are supported."))

    if not source_segments:
        raise ValueError(L('validation_7', "No non-empty source segments were found."))

    if not target_segments:
        raise ValueError(L('validation_8', "No non-empty target segments were found."))

    metadata["source_count"] = len(source_segments)
    metadata["target_count"] = len(target_segments)

    if len(source_segments) != len(target_segments):
        raise ValueError(
            (LANGUAGE["ko"]['validation_9'].format(source_count=len(source_segments), target_count=len(target_segments)) if UI_LANGUAGE == "ko" else "Document structure mismatch: "
            f"source has {len(source_segments)} {metadata['unit']}, "
            f"but target has {len(target_segments)}. "
            "Align the documents before running QA.")
        )

    dataframe = pd.DataFrame(
        {
            "Segment": range(1, len(source_segments) + 1),
            "Source": source_segments,
            "Target": target_segments,
        }
    )
    return dataframe, metadata


def compare_items(source_items, target_items):
    return Counter(source_items) != Counter(target_items)


def normalize_term(text):
    """용어 비교를 위해 공백과 영문 대소문자를 정규화합니다."""
    if pd.isna(text):
        return ""

    normalized = re.sub(r"\s+", " ", str(text).strip())
    return normalized.casefold()


def load_glossary(glossary_file):
    """XLSX/CSV 용어집을 읽고 원문·번역문 열 이름의 표기 차이를 허용합니다."""
    if glossary_file is None:
        return []

    file_name = glossary_file.name.lower()

    try:
        if file_name.endswith(".xlsx"):
            glossary_dataframe = pd.read_excel(glossary_file)
        elif file_name.endswith(".csv"):
            try:
                glossary_dataframe = pd.read_csv(
                    glossary_file,
                    encoding="utf-8-sig"
                )
            except UnicodeDecodeError:
                glossary_file.seek(0)
                glossary_dataframe = pd.read_csv(
                    glossary_file,
                    encoding="cp949"
                )
        else:
            raise ValueError(L('validation_10', "Glossary file must be XLSX or CSV."))
    except Exception as error:
        raise ValueError(
            L('validation_11', f"Failed to read glossary file: {error}")
        ) from error

    glossary_dataframe.columns = [str(column).strip() for column in glossary_dataframe.columns]

    # "Source Term"과 "source_term"처럼 공백·밑줄·대소문자만 다른
    # 열 이름을 동일하게 취급합니다. 기존 회귀 용어집과 일반 운영용
    # 용어집을 모두 지원하기 위한 하위 호환 규칙입니다.
    normalized_columns = {
        re.sub(r"[\s_-]+", "", column).casefold(): column
        for column in glossary_dataframe.columns
    }
    source_column = normalized_columns.get("sourceterm")
    target_column = normalized_columns.get("targetterm")

    if source_column is None or target_column is None:
        raise ValueError(
            L('validation_12', "Glossary file must contain source and target term columns. "
            "Supported names: Source Term/Target Term or "
            "source_term/target_term.")
        )

    glossary_entries = []

    for _, row in glossary_dataframe.iterrows():
        source_term = normalize_text(row[source_column])
        target_term_text = normalize_text(row[target_column])

        if not source_term or not target_term_text:
            continue

        approved_targets = [
            target.strip()
            for target in target_term_text.split("|")
            if target.strip()
        ]

        if approved_targets:
            glossary_entries.append(
                {
                    "source_term": source_term,
                    "approved_targets": approved_targets
                }
            )

    return glossary_entries


def contains_source_term(source_text, source_term):
    """원문에서 용어를 검색하며 영문 용어에는 단어 경계를 적용합니다."""
    normalized_source = normalize_term(source_text)
    normalized_term = normalize_term(source_term)

    if not normalized_term:
        return False

    is_word_based = bool(
        re.fullmatch(r"[A-Za-z0-9 _.-]+", source_term)
    )

    if is_word_based:
        pattern = (
            r"(?<!\w)"
            + re.escape(normalized_term)
            + r"(?!\w)"
        )
        return bool(re.search(pattern, normalized_source))

    return normalized_term in normalized_source


def contains_approved_target(target_text, approved_targets):
    """번역문에 승인 번역어 중 하나가 포함되어 있는지 확인합니다."""
    normalized_target = normalize_term(target_text)

    return any(
        normalize_term(approved_target) in normalized_target
        for approved_target in approved_targets
        if normalize_term(approved_target)
    )


def check_terminology(source_text, target_text, glossary_entries):
    """원문 용어가 있을 때 승인 번역어 사용 여부를 검사합니다."""
    terminology_issues = []

    for entry in glossary_entries or []:
        source_term = entry["source_term"]
        approved_targets = entry["approved_targets"]

        if not contains_source_term(source_text, source_term):
            continue

        if contains_approved_target(target_text, approved_targets):
            continue

        terminology_issues.append(
            f"{source_term} → {' / '.join(approved_targets)}"
        )

    return terminology_issues


# --------------------------------------------------
# 숫자 검사
# --------------------------------------------------
def extract_numbers(text):
    text = normalize_text(text)

    pattern = r"\d+(?:,\d{3})*(?:\.\d+)?"

    return re.findall(pattern, text)


# --------------------------------------------------
# 문장 끝 구두점 검사
# --------------------------------------------------
def extract_end_punctuation(text):
    text = normalize_text(text)

    if not text:
        return ""

    match = re.search(r"[.!?:;]+$", text)

    if match:
        return match.group()

    return ""


# --------------------------------------------------
# 태그 검사
# --------------------------------------------------
def extract_tags(text):
    text = normalize_text(text)

    pattern = r"</?[A-Za-z][^>]*?>"

    return re.findall(pattern, text)


# --------------------------------------------------
# 자리표시자 검사
# --------------------------------------------------
def extract_placeholders(text):
    text = normalize_text(text)

    patterns = [
        r"\{[^{}]+\}",          # {0}, {name}
        r"%[-+#0 ]*\d*(?:\.\d+)?[a-zA-Z]",  # %s, %d, %02d
        r"\$\{[^{}]+\}",        # ${name}
        r"\{\{\s*[^{}]+\s*\}\}",  # {{name}}
    ]

    placeholders = []

    for pattern in patterns:
        placeholders.extend(re.findall(pattern, text))

    return placeholders


# --------------------------------------------------
# 괄호 검사
# --------------------------------------------------
def extract_brackets(text):
    text = normalize_text(text)

    return re.findall(r"[\(\)\[\]\{\}]", text)


def check_balanced_brackets(text):
    text = normalize_text(text)

    bracket_pairs = {
        ")": "(",
        "]": "[",
        "}": "{"
    }

    opening_brackets = set(bracket_pairs.values())
    closing_brackets = set(bracket_pairs.keys())

    stack = []

    for character in text:
        if character in opening_brackets:
            stack.append(character)

        elif character in closing_brackets:
            if not stack:
                return False

            if stack[-1] != bracket_pairs[character]:
                return False

            stack.pop()

    return len(stack) == 0


# --------------------------------------------------
# URL 검사
# --------------------------------------------------
def extract_urls(text):
    text = normalize_text(text)

    pattern = r"https?://[^\s<>\"]+"

    return re.findall(pattern, text)


# --------------------------------------------------
# 이메일 주소 검사
# --------------------------------------------------
def extract_emails(text):
    text = normalize_text(text)

    pattern = r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"

    return re.findall(pattern, text)


# --------------------------------------------------
# 공백 검사
# --------------------------------------------------
def check_whitespace_issues(text):
    """
    Target 문장 내부의 비정상적인 공백을 검사합니다.

    검사 대상:
    - 문장 앞 공백
    - 문장 뒤 공백
    - 두 칸 이상의 연속 공백
    - 탭 문자
    """
    if pd.isna(text):
        return []

    original_text = str(text)
    issues = []

    if original_text.startswith(" "):
        issues.append("Leading space")

    if original_text.endswith(" "):
        issues.append("Trailing space")

    if re.search(r" {2,}", original_text):
        issues.append("Multiple spaces")

    if "\t" in original_text:
        issues.append("Tab character")

    return issues


# --------------------------------------------------
# 파일명 검사
# --------------------------------------------------
def extract_file_names(text):
    """
    확장자를 포함한 파일명을 추출합니다.

    예:
    config.yaml
    README.md
    install.exe
    archive.tar.gz
    """
    text = normalize_text(text)

    pattern = (
        r"\b"
        r"[A-Za-z0-9_-]+"
        r"(?:\.[A-Za-z0-9_-]+)+"
        r"\b"
    )

    return re.findall(pattern, text)


# --------------------------------------------------
# 키보드 단축키 검사
# --------------------------------------------------
def extract_shortcuts(text):
    """
    키보드 단축키를 추출합니다.

    예:
    Ctrl+C
    Ctrl + C
    Alt+F4
    Shift+Enter
    F5
    """
    text = normalize_text(text)

    pattern = (
        r"\b(?:Ctrl|Alt|Shift|Command|Cmd|Option|Windows|Win)"
        r"\s*\+\s*"
        r"(?:[A-Za-z0-9]+|F\d{1,2})\b"
        r"|"
        r"\bF(?:[1-9]|1[0-2])\b"
    )

    shortcuts = re.findall(
        pattern,
        text,
        flags=re.IGNORECASE
    )

    # Ctrl + C와 Ctrl+C를 동일하게 비교하기 위한 정규화
    normalized_shortcuts = []

    for shortcut in shortcuts:
        normalized = re.sub(
            r"\s+",
            "",
            shortcut
        ).upper()

        normalized_shortcuts.append(normalized)

    return normalized_shortcuts


# --------------------------------------------------
# 단위 검사
# --------------------------------------------------
def extract_units(text):
    if not isinstance(text, str):
        return []

    units = [
        "mm", "cm", "m", "km",
        "mg", "g", "kg",
        "ml", "l",
        "ms", "sec", "second", "seconds",
        "min", "minute", "minutes",
        "hour", "hours",
        "kb", "mb", "gb", "tb",
        "hz", "khz", "mhz", "ghz",
        "v", "w", "kw",
        "%", "°c", "°f"
    ]

    # 정규식 특수문자를 안전하게 처리
    escaped_units = [re.escape(unit) for unit in units]

    unit_pattern = (
        r"\b\d+(?:[.,]\d+)?\s*(?:"
        + "|".join(escaped_units)
        + r")\b"
    )

    return re.findall(
        unit_pattern,
        text,
        flags=re.IGNORECASE
    )


# --------------------------------------------------
# QA 규칙
# --------------------------------------------------
def run_qa_check(source, target, glossary_entries=None):
    if glossary_entries is None:
        glossary_entries = []

    source_text = normalize_text(source)
    target_text = normalize_text(target)

    issues = []

    # 번역 누락
    if not target_text:
        issues.append("Missing translation")
        return issues

    # 원문과 번역문 동일
    if source_text == target_text:
        issues.append(
            "Source and target are identical"
        )

    # 숫자
    if compare_items(
        extract_numbers(source_text),
        extract_numbers(target_text)
    ):
        issues.append("Number mismatch")

    # 문장 끝 구두점
    if (
        extract_end_punctuation(source_text)
        != extract_end_punctuation(target_text)
    ):
        issues.append("Punctuation mismatch")

    # HTML/XML 태그
    if compare_items(
        extract_tags(source_text),
        extract_tags(target_text)
    ):
        issues.append("Tag mismatch")

    # 자리표시자
    if compare_items(
        extract_placeholders(source_text),
        extract_placeholders(target_text)
    ):
        issues.append("Placeholder mismatch")

    # 괄호 문자
    if compare_items(
        extract_brackets(source_text),
        extract_brackets(target_text)
    ):
        issues.append("Bracket mismatch")

    # Target 내부 괄호 균형
    if not check_balanced_brackets(target_text):
        issues.append(
            "Unbalanced target brackets"
        )

    # URL
    if compare_items(
        extract_urls(source_text),
        extract_urls(target_text)
    ):
        issues.append("URL mismatch")

    # 이메일
    if compare_items(
        extract_emails(source_text),
        extract_emails(target_text)
    ):
        issues.append("Email mismatch")

    # 공백
    whitespace_issues = check_whitespace_issues(
        target
    )

    if whitespace_issues:
        issues.append(
            "Whitespace warning: "
            + ", ".join(whitespace_issues)
        )

    # 파일명
    if compare_items(
        extract_file_names(source_text),
        extract_file_names(target_text)
    ):
        issues.append("File name mismatch")

    # 키보드 단축키
    if compare_items(
        extract_shortcuts(source_text),
        extract_shortcuts(target_text)
    ):
        issues.append("Shortcut mismatch")

    # 단위
    if compare_items(
        extract_units(source_text),
        extract_units(target_text)
    ):
        issues.append("Unit mismatch")

    # 용어집 기반 Terminology QA
    terminology_issues = check_terminology(
        source_text,
        target_text,
        glossary_entries
    )

    for terminology_issue in terminology_issues:
        issues.append(
            "Terminology mismatch: "
            + terminology_issue
        )

    return issues


def convert_dataframe_to_excel(dataframe):
    output = io.BytesIO()

    # Prevent uploaded text from being interpreted as a spreadsheet formula
    # when a user opens the downloaded QA report.
    safe_dataframe = dataframe.map(
        lambda value: "'" + value
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
        else value
    )

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:
        safe_dataframe.to_excel(
            writer,
            index=False,
            sheet_name="QA Results"
        )

    return output.getvalue()


# --------------------------------------------------
# 파일 업로드
# --------------------------------------------------
st.markdown(f'<div class="ltli-section">{L("add_files", "1. Add files")}</div>', unsafe_allow_html=True)
st.caption(L("choose_format", "Choose the input format that matches your translation workflow."))

# Download examples before uploading files. XLSX uses Source/Target columns;
# DOCX requires a separate source and target document.
downloads_dir = Path(__file__).resolve().parent / "qa_downloads"

with st.expander(
    L("downloads_title", "📥 Download sample files and blank templates"),
    expanded=False,
):
    st.caption(
        L(
            "downloads_help",
            "The Excel sample is one bilingual file. For Word, download both source and target files and upload them together. Samples contain fictional data and intentional QA issues.",
        )
    )
    for item_type, heading in [
        ("sample", L("sample_label", "Sample files · ready to test")),
        ("template", L("template_label", "Blank templates · add your own content")),
    ]:
        st.markdown(f"**{heading}**")
        excel_col, source_col, target_col = st.columns(3)
        items = [
            (excel_col, f"LTLI_QA_{item_type}.xlsx", L("download_xlsx", "⬇ Excel (.xlsx)")),
            (source_col, f"LTLI_QA_{item_type}_source.docx", L("download_source", "⬇ Word source (.docx)")),
            (target_col, f"LTLI_QA_{item_type}_target.docx", L("download_target", "⬇ Word target (.docx)")),
        ]
        for column, filename, label in items:
            filepath = downloads_dir / filename
            with column:
                if filepath.is_file():
                    st.download_button(
                        label=label,
                        data=filepath.read_bytes(),
                        file_name=filename,
                        mime=(
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                            if filename.endswith(".xlsx")
                            else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        ),
                        key=f"qa_download_{item_type}_{filename}",
                        use_container_width=True,
                    )
                else:
                    st.caption(
                        L("downloads_unavailable", "File Not Available")
                    )
    st.caption(
        L(
            "template_help",
            "Excel templates contain Source and Target headers. Word templates are separate empty documents; keep paragraph alignment between source and target.",
        )
    )

xlsx_tab, paired_tab = st.tabs([L("xlsx_tab", "📊 Bilingual XLSX"), L("pair_tab", "📄 Separate DOCX / TXT")])

uploaded_file = None
source_file = None
target_file = None

with xlsx_tab:
    uploaded_file = st.file_uploader(
        L("drop_xlsx", "Drop a bilingual spreadsheet here"),
        type=["xlsx"],
        help=L('xlsx_help', 'The spreadsheet needs exactly one Source column and one Target column.'),
        key="bilingual_xlsx_uploader",
    )
    st.caption(L("required_columns", "Required columns: Source and Target · Maximum 10 MB"))

with paired_tab:
    source_column, target_column = st.columns(2)
    with source_column:
        source_file = st.file_uploader(
            L("source_file", "Source file"), type=["docx", "txt"], key="source_document_uploader"
        )
    with target_column:
        target_file = st.file_uploader(
            L("target_file", "Target file"), type=["docx", "txt"], key="target_document_uploader"
        )
    st.caption(
        L("pair_help", "Use two files of the same type. DOCX is aligned by non-empty paragraph; TXT is aligned by non-empty line.")
    )

if uploaded_file is not None:
    input_mode = "Bilingual XLSX"
elif source_file is not None or target_file is not None:
    input_mode = "Separate source and target files"
else:
    input_mode = None

with st.expander(L("glossary", "＋ Add a terminology glossary (optional)")):
    glossary_file = st.file_uploader(
        L("glossary_file", "Glossary file"), type=["csv", "xlsx"], key="glossary_uploader",
        help=L('glossary_help', 'Columns: Source Term and Target Term (space, underscore, and case variations accepted).')
    )

glossary_entries = []

if glossary_file is not None:
    try:
        glossary_entries = load_glossary(glossary_file)
        st.success(
            L('glossary_loaded', 'Loaded {count} glossary entries.').format(count=len(glossary_entries))
        )
    except ValueError as error:
        st.error(str(error))
        st.stop()


alignment_metadata = None

if input_mode is None:
    st.info(L("begin", "Choose a tab above and upload your file(s) to begin."), icon="👆")
    st.stop()
elif input_mode == "Bilingual XLSX":
    if uploaded_file is None:
        st.info(L('upload_xlsx', 'Please upload a bilingual XLSX file.'))
        st.stop()

    try:
        validate_file_size(uploaded_file)
        dataframe = pd.read_excel(uploaded_file)
    except Exception as error:
        track_event("qa_failed", {"stage": "file_read", "error_type": type(error).__name__})
        st.error(L('file_read_failed', 'Failed to read the file. Check the format and file size (up to 10 MB).'))
        st.stop()

    st.success(L('uploaded', 'Uploaded: {name}').format(name=uploaded_file.name))
else:
    if source_file is None or target_file is None:
        st.info(L('upload_pair', 'Please upload both the source file and the target file.'))
        st.stop()

    try:
        dataframe, alignment_metadata = build_parallel_dataframe(
            source_file,
            target_file,
        )
    except ValueError as error:
        track_event("qa_failed", {"stage": "alignment", "error_type": type(error).__name__})
        st.error(str(error))
        st.stop()

    st.success(
        L('aligned', 'Aligned {count} {unit}.').format(count=alignment_metadata['source_count'], unit=L(alignment_metadata['unit'], alignment_metadata['unit']))
    )

    metric1, metric2, metric3 = st.columns(3)
    metric1.metric(L('format', 'Format'), alignment_metadata["format"])
    metric2.metric(L('source_units', 'Source units'), alignment_metadata["source_count"])
    metric3.metric(L('target_units', 'Target units'), alignment_metadata["target_count"])


# --------------------------------------------------
# 열 확인 및 표준화
# --------------------------------------------------
try:
    dataframe = normalize_bilingual_columns(dataframe)
except ValueError as error:
    track_event("qa_failed", {"stage": "column_validation", "error_type": type(error).__name__})
    st.error(
        str(error)
    )

    st.write(L('current_columns', 'Current columns:'))
    st.write(list(dataframe.columns))

    st.stop()

if input_mode == "Bilingual XLSX":
    analytics_files = [uploaded_file]
else:
    analytics_files = [source_file, target_file]

file_formats = [get_file_extension(item).upper() for item in analytics_files]
size_buckets = [file_size_bucket(getattr(item, "size", None)) for item in analytics_files]
upload_event_key = "|".join([input_mode, *file_formats, *size_buckets])
if st.session_state.get("active_upload_key") != upload_event_key:
    for result_key in ("qa_result_dataframe", "qa_error_dataframe", "qa_issue_summary"):
        st.session_state.pop(result_key, None)
    st.session_state["active_upload_key"] = upload_event_key
track_once(
    f"file_uploaded_{upload_event_key}",
    "file_uploaded",
    {"file_formats": file_formats, "file_count": len(analytics_files), "size_buckets": size_buckets},
)


# --------------------------------------------------
# 파일 미리보기
# --------------------------------------------------
st.markdown(f'<div class="ltli-section">{L("review", "2. Review and run QA")}</div>', unsafe_allow_html=True)
st.caption(f"{len(dataframe):,}{L('ready', ' aligned segment(s) are ready. Confirm the preview, then start the check.')}")

st.dataframe(
    localized_results(dataframe),
    use_container_width=True,
    hide_index=True
)

run_column, detail_column = st.columns([1, 2])
with run_column:
    run_requested = st.button(L("run", "Run QA checks  →"), type="primary", use_container_width=True)
with detail_column:
    st.caption(L('checks_description', 'Checks missing text, identical segments, numbers, punctuation, tags, placeholders, brackets, URLs, emails, whitespace, filenames, shortcuts, units, and terminology. Automated QA does not replace final human review.'))

# --------------------------------------------------
# QA 실행
# --------------------------------------------------
issue_types = [
    "Missing translation",
    "Source and target are identical",
    "Number mismatch",
    "Punctuation mismatch",
    "Tag mismatch",
    "Placeholder mismatch",
    "Bracket mismatch",
    "Unbalanced target brackets",
    "URL mismatch",
    "Email mismatch",
    "Whitespace warning",
    "File name mismatch",
    "Shortcut mismatch",
    "Unit mismatch",
    "Terminology mismatch"
]

if run_requested:
    qa_started_at = time.perf_counter()
    qa_metadata = {"input_mode": input_mode, "glossary_used": bool(glossary_entries), "segment_count": len(dataframe)}
    track_event("qa_started", qa_metadata)
    try:
        result_dataframe = dataframe.copy()
        result_dataframe["QA Issues"] = result_dataframe.apply(
            lambda row: run_qa_check(row["Source"], row["Target"], glossary_entries), axis=1
        )
        result_dataframe["QA Result"] = result_dataframe["QA Issues"].apply(
            lambda issues: "Pass" if not issues else "Error"
        )
        result_dataframe["QA Issues"] = result_dataframe["QA Issues"].apply(
            lambda issues: ", ".join(issues)
        )
        error_dataframe = result_dataframe[result_dataframe["QA Result"] == "Error"].copy()
        issue_summary = pd.DataFrame({
            "Issue Type": issue_types,
            "Count": [int(result_dataframe["QA Issues"].str.contains(issue, regex=False).sum()) for issue in issue_types],
        })
        st.session_state["qa_result_dataframe"] = result_dataframe
        st.session_state["qa_error_dataframe"] = error_dataframe
        st.session_state["qa_issue_summary"] = issue_summary
        track_event("qa_completed", {
            **qa_metadata,
            "processing_ms": elapsed_milliseconds(qa_started_at),
            "error_segment_count": len(error_dataframe),
            "issue_counts": {
                str(row["Issue Type"]): int(row["Count"])
                for _, row in issue_summary.iterrows()
            },
        })
    except Exception as error:
        track_event("qa_failed", {"stage": "qa_processing", "error_type": type(error).__name__})
        st.error(L('qa_failed', 'QA could not be completed. Please check the input and try again.'))

if "qa_result_dataframe" in st.session_state:
    result_dataframe = st.session_state["qa_result_dataframe"]
    error_dataframe = st.session_state["qa_error_dataframe"]
    issue_summary = st.session_state["qa_issue_summary"]
    total_segments = len(result_dataframe)
    error_count = len(error_dataframe)
    pass_count = total_segments - error_count
    pass_rate = pass_count / total_segments * 100 if total_segments else 0

    st.markdown(f'<div class="ltli-section">{L("results", "3. QA results")}</div>', unsafe_allow_html=True)
    column1, column2, column3, column4 = st.columns(4)
    column1.metric(L('total', 'Total Segments'), total_segments)
    column2.metric(L('passed', 'Passed'), pass_count)
    column3.metric(L('errors', 'Issues'), error_count)
    column4.metric(L('pass_rate', 'Pass Rate'), f"{pass_rate:.1f}%")
    st.markdown(L("overview", "#### Issue overview"))
    st.dataframe(localized_results(issue_summary), use_container_width=True, hide_index=True)
    st.markdown(L("all_segments", "#### All segments"))
    st.dataframe(localized_results(result_dataframe), use_container_width=True, hide_index=True)
    if error_count:
        st.error(L('issue_count', '{count} segment(s) contain QA issues.').format(count=error_count))
        st.markdown(L("review_issues", "#### Issues requiring review"))
        st.dataframe(localized_results(error_dataframe), use_container_width=True, hide_index=True)
    else:
        st.success(L("no_issues", "No QA issues were found."))

    def record_report_download():
        track_event("report_downloaded", {"report_format": "XLSX", "segment_count": total_segments})

    st.download_button(
        label=L("download", "Download Excel report  ↓"),
        data=convert_dataframe_to_excel(result_dataframe),
        file_name="ltli_qa_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        on_click=record_report_download,
        use_container_width=True,
    )


st.markdown(
    '<div class="ltli-footer">' + L("footer", "Language Technology & Localization Institute · LTLI QA · Built for translators, reviewers, and localization teams") + ' · <a href="https://ltli.ai">LTLI</a> · <a href="mailto:contact@ltli.kr">contact@ltli.kr</a></div>',
    unsafe_allow_html=True,
)
